import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

products = json.load(open(sys.argv[1]))
OUT = sys.argv[2]

FOREST, GOLD, CREAM = "1B4332", "FDC717", "FAF6E9"
H_FILL = PatternFill("solid", fgColor=FOREST)
EXIST = PatternFill("solid", fgColor="E8EFE9")   # pre-filled from old site
TOFILL = PatternFill("solid", fgColor="FFF4C2")  # Hasiru team to fill
CALC = PatternFill("solid", fgColor="DCE6F2")    # formula, do not edit
SEC = PatternFill("solid", fgColor=CREAM)
thin = Side(style="thin", color="C9C9C9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HFONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BFONT = Font(name="Arial", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.border = H_FILL, HFONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_body(ws, r0, r1, ncols, fills=None):
    for r in range(r0, r1 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border, cell.alignment = BFONT, BORDER, WRAP
            if fills:
                cell.fill = fills[c - 1]


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color=FOREST)
    ws["A2"] = sub
    ws["A2"].font = Font(name="Arial", italic=True, size=10, color="5A5A5A")


def dv_list(ws, items, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


# ---------- 1. Read Me ----------
ws = wb.active
ws.title = "Read Me"
title(ws, "Hasiru Agro: New E-Commerce Website, Data Collection Sheet",
      "Prepared by WODO for the Hasiru Agro team. Please fill the yellow cells in every tab.")
rows = [
    ("Purpose", "The new website needs product, recommendation, calculator and delivery data to power the features in the SRS. This workbook lists every feature, the data it needs, and gives a place to fill it."),
    ("Colour key", ""),
    ("", "Green cells: already taken from the old hasiruagro.com shop. Please check and correct."),
    ("", "Yellow cells: please fill. These are new and required for the new features."),
    ("", "Blue cells: calculated automatically. Do not edit."),
    ("Tabs", ""),
    ("1. Feature Map", "Every website feature, the SRS section it comes from, and which tab/columns feed it."),
    ("2. Plants", "All 193 plants from the old shop, pre-filled. Add new columns (category, weight, spacing, model role, etc.). Add new plants at the bottom."),
    ("3. Input Products", "Compost, vermicompost, biochar, seeds, liquid fertiliser, bio-pesticide. New, please fill fully."),
    ("4. Recommendation Models", "Commercial, Food Forest, Miyawaki: which plants, their role (main/host/biomass), spacing and quantities."),
    ("5. Calculators", "Formulas for the plant, compost, biochar, vermicompost and seed calculators."),
    ("6. Logistics", "Vehicles, capacity, delivery zones, Bengaluru schedule, outstation rates, packing charges."),
    ("7. Content & Policies", "Checklist of text, logos, photos, videos, testimonials, client and project data."),
    ("8. Category Map", "Proposed mapping of old shop categories to the new Timber / Fruiting / Other structure."),
    ("Notes", ""),
    ("Pricing", "Each plant has a normal price (below its bulk quantity) and a bulk price. The bulk quantity differs per plant: e.g. 50 for big trees like mango, 200 for small shrubs. Fill both columns in the Plants tab. Old Regular / Sale / Offline prices are shown for reference."),
    ("Sizes", "Plant sizes: do any plants come in more than one size (e.g. 1.5 ft and 6 ft)? If yes, add one row per size and use the same Parent Name."),
    ("Weights", "Weights: every plant currently shows 1 kg / 70 cm (placeholder). Real weight and pot/bag size are needed for delivery and truck-capacity calculation."),
    ("GST", "GST: please confirm HSN and GST % per product (plants are usually HSN 0602 at 0%, vermicompost/compost 3101 at 5%; confirm with your CA)."),
]
for i, (a, b) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=a).font = Font(name="Arial", bold=True, size=10, color=FOREST)
    c = ws.cell(row=i, column=2, value=b)
    c.font, c.alignment = BFONT, WRAP
for r, f in ((6, EXIST), (7, TOFILL), (8, CALC)):
    ws.cell(row=r, column=2).fill = f
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 120

# ---------- 2. Feature Map ----------
ws = wb.create_sheet("Feature Map")
title(ws, "Feature Map", "Each website feature and the data needed to build it.")
cols = ["#", "Feature", "Where on site", "SRS ref", "What it does", "Data needed from Hasiru Agro", "Filled in tab", "Priority", "Status"]
header(ws, 4, cols, [5, 28, 22, 9, 48, 52, 22, 11, 14])
features = [
    ("Two-path entry (I know what I want / Help me choose)", "Home hero", "1.1", "Splits shoppers into direct shop vs guided recommendation.", "Banner image, headline copy.", "Content & Policies", "Must"),
    ("Best sellers carousel (quick-commerce style)", "Home", "1.2", "Horizontal swipe cards with add to cart.", "Best seller flag, image, height, exposure, rain, temperature, packet weight, price, offer price, availability.", "Plants", "Must"),
    ("Simple plant calculator (plants per acre)", "Home + floating on Shop", "1.2A, 3.4", "43,560 / (spacing L x W ft) = plants per acre.", "Recommended spacing L and W (ft) per plant. Confirm formula and any buffer % (note '43,560 x 3 / 100' in sketch).", "Plants, Calculators", "Must"),
    ("Input products carousel", "Home", "1.3", "Compost, seeds, biochar, liquid fertiliser, bio-pesticide cards.", "Full input product list with pack sizes and prices.", "Input Products", "Must"),
    ("Compost calculator", "Home", "1.4", "Plants selected -> estimated compost -> recommended product -> add to cart.", "Compost kg per plant (by plant size/type), product to recommend.", "Calculators, Plants", "Must"),
    ("Incubated by UAS Dharwad band", "Home + About", "1.5, 26", "Running band of institutional logos.", "Logos (UAS Dharwad, Startup India, MSME, others), approved wording, links.", "Content & Policies", "Must"),
    ("DIY / Learn how videos", "Home + Recommendation page", "1.6, 8", "Video carousel; model-specific planting video.", "Video URLs, titles, thumbnails, category, linked model.", "Content & Policies", "Must"),
    ("Why Choose Us", "Home + own page", "1.7, 2, 24", "Authenticity, expertise, results, delivery to farm.", "Verified numbers: acres, farms, customers, projects, plants planted.", "Content & Policies", "Must"),
    ("Testimonials", "Home", "1.8", "Horizontal review cards.", "Name, location, photo, review, rating, farm/project info.", "Content & Policies", "Must"),
    ("Plant shop with categories", "Shop", "3.1-3.3", "Timber / Fruiting / Other grid with qty +/-.", "Final category per plant, description, price, stock, availability.", "Plants, Category Map", "Must"),
    ("Product attributes", "Product page", "1.2, 3.2", "Spec table on each plant.", "Height, exposure, rain, temperature, type, bearing time, yield, spacing, harvest season.", "Plants", "Must"),
    ("Help me choose wizard", "Popup", "4.1-4.3", "Land size (acre / sq ft), location (map/search), purpose, name + phone OTP, Generate.", "Purposes list (Commercial, Food Forest, Miyawaki, others). Any location-based rules (which plants suit which district/climate).", "Recommendation Models", "Must"),
    ("Recommendation page (3 models)", "Recommendation", "5.1-5.6", "All 3 models generated, expandable, each with plants by category and recommended quantity.", "Per model: plants, category, role, spacing, quantity per 36x36 ft and per acre.", "Recommendation Models", "Must"),
    ("36 x 36 ft plantation diagram", "Recommendation", "5.3", "Visual layout of main / host / biomass plants and spacing.", "Planting pattern for each model (grid drawing or row pattern like M o X o X o M), legend.", "Recommendation Models", "Must"),
    ("Recommended inputs (biochar, vermicompost)", "Recommendation + Cart", "6, 6.1", "Calculates input kg from plants selected; suggests products.", "kg per plant for each input, by plant type/size and model.", "Calculators, Plants", "Must"),
    ("Recommended package in cart", "Cart", "9", "Recommended items in a highlighted block; extra items below.", "None beyond models + products.", "-", "Must"),
    ("Truck capacity meter", "Cart", "13", "'Truck 65% full, add ~35% more'.", "Vehicle types, capacity (plants / kg / volume), weight and pot size per product.", "Logistics, Plants", "Must"),
    ("Delivery options", "Checkout", "12, 14, 15, 29", "Nursery pickup (free), Bengaluru weekly slot, outstation by km + weight.", "Bengaluru delivery day/slot, zones/PIN codes, per-km rates, minimum charges, outstation rules.", "Logistics", "Must"),
    ("Packing charges", "Checkout", "16", "Per order / product / qty / size based.", "Packing rule and rates.", "Logistics", "Must"),
    ("GST calculation + invoices", "Checkout", "11, 41", "Per product tax; CGST/SGST/IGST; B2B GSTIN.", "HSN code and GST % per product; company GSTIN and invoice details.", "Plants, Input Products", "Must"),
    ("Offers and coupons", "Cart / Checkout", "18, 33", "Admin-controlled % / fixed discounts, campaigns, coupons.", "Current/planned offers (e.g. SUMMER15 on old site).", "Content & Policies", "Should"),
    ("Policies", "Checkout + footer", "19, 20", "No cancellation / return; damage and replacement policy.", "Final policy wording (damage reporting window, proof needed, replacement process).", "Content & Policies", "Must"),
    ("Farm development services", "Farm Development page", "21-25", "Assessment, design, water, earthwork, planting; Our Work; Our Clients; enquiry form.", "Project list (images, location, acres, scope, before/after), client list.", "Content & Policies", "Must"),
    ("About Us", "About page", "27", "Chemical free, certified permaculture designer, vision, mission.", "Certificates, vision, mission text.", "Content & Policies", "Must"),
    ("Customer accounts + saved land profiles", "Account", "38", "OTP login, order history, invoices, saved land profiles.", "SMS / WhatsApp provider account (for OTP).", "Content & Policies", "Should"),
    ("English / Kannada", "Whole site", "39", "Language toggle.", "Kannada name for each plant and category; Kannada translations of site text.", "Plants, Content & Policies", "Should"),
    ("WhatsApp / SMS alerts", "Backend", "40", "OTP, order confirmations, delivery reminders, enquiry alerts.", "WhatsApp Business number, SMS provider, admin phone numbers for alerts.", "Content & Policies", "Should"),
    ("Payment gateway", "Checkout", "41", "UPI, cards, net banking, wallets.", "Razorpay / PhonePe / Cashfree account (KYC done).", "Content & Policies", "Must"),
    ("Admin dashboard + inventory", "Admin", "30-33", "Sales, regions, best sellers, low stock, enquiries, orders.", "Low-stock threshold per product; current stock.", "Plants, Input Products", "Must"),
]
for i, f in enumerate(features, 1):
    r = 4 + i
    for c, v in enumerate((i,) + f + ("Pending",), 1):
        ws.cell(row=r, column=c, value=v)
style_body(ws, 5, 4 + len(features), len(cols))
dv_list(ws, ["Must", "Should", "Later"], f"H5:H{4+len(features)}")
dv_list(ws, ["Pending", "In progress", "Data received", "Done"], f"I5:I{4+len(features)}")

# ---------- 3. Plants ----------
ws = wb.create_sheet("Plants")
title(ws, "Plants: Product Master", "Green = from old shop (please verify). Yellow = please fill. Blue = auto. One row per plant per size.")
OLD2NEW = {"Forestry": ("Timber Trees", ""), "Desi Fruits": ("Fruiting Trees", "Desi"),
           "Exotic fruits": ("Fruiting Trees", "Exotic"), "Spices": ("Other Plants", "Spices"),
           "Plantation": ("Other Plants", "Plantation")}
spec = [  # (header, width, kind)  kind: E existing, F fill, C calc
    ("SKU", 10, "E"), ("Plant Name", 26, "E"), ("Kannada Name", 18, "F"), ("Old Category", 14, "E"),
    ("New Category", 15, "F"), ("Sub-category", 12, "F"), ("Show on site?", 10, "F"), ("Best Seller?", 10, "F"),
    ("Plant Type (Grafted / Seeded / Air layered...)", 14, "E"), ("Size / Variant (e.g. 2 ft, 6 ft)", 14, "F"),
    ("Plant Height at Sale (ft)", 11, "F"), ("Pot / Bag Size", 12, "F"), ("Weight per Plant incl. pot (kg)", 12, "F"),
    ("Plants per Truck (smallest vehicle)", 12, "F"),
    ("Regular Price (Rs)", 10, "E"), ("Sale Price (Rs)", 10, "E"), ("Offline Price (Rs)", 10, "E"),
    ("Bulk Starts at Qty (e.g. 50 for big trees, 200 for shrubs)", 13, "F"), ("Bulk Price per Plant (Rs)", 10, "F"), ("Current Stock", 9, "E"), ("Low Stock Alert at", 9, "F"),
    ("Availability", 14, "F"), ("HSN Code", 9, "F"), ("GST %", 7, "F"),
    ("Spacing L (ft)", 8, "E"), ("Spacing W (ft)", 8, "E"), ("Plants per Acre", 9, "C"),
    ("Spacing (old text)", 14, "E"), ("Exposure", 14, "E"), ("Rain / Water Need", 12, "F"), ("Temperature", 14, "E"),
    ("Growth Habit", 16, "E"), ("Bearing Time", 10, "E"), ("Yield per Tree", 16, "E"), ("Harvest Season", 16, "E"),
    ("Fruit Size", 12, "E"), ("Suitable Models (Commercial / Food Forest / Miyawaki)", 18, "F"),
    ("Role in Model (Main / Host / Biomass)", 12, "F"), ("Forest Layer (Canopy / Sub-canopy / Shrub...)", 13, "F"),
    ("Vermicompost per Plant (kg)", 11, "F"), ("Biochar per Plant (kg)", 10, "F"), ("Compost per Plant (kg)", 10, "F"),
    ("Districts / Zones NOT suitable", 16, "F"), ("Image URL (old site)", 30, "E"), ("New Photo Needed?", 10, "F"),
    ("Notes / Corrections", 24, "F"),
]
header(ws, 4, [s[0] for s in spec], [s[1] for s in spec])
fills = [{"E": EXIST, "F": TOFILL, "C": CALC}[s[2]] for s in spec]
col = {s[0]: i for i, s in enumerate(spec, 1)}


def num(v):
    try:
        return float(v) if "." in v else int(v)
    except (ValueError, TypeError):
        return None


TYPEFIX = {"gafted": "Grafted", "air layered": "Air layered", "grafted": "Grafted", "seeded": "Seeded"}
r = 5
for p in products:
    if p["Type"] != "simple" or "Uncategorized" in p["Categories"]:
        continue  # skips the e-book
    oldcat = p["Categories"].split(",")[0].strip()
    newcat, sub = OLD2NEW.get(oldcat, ("", ""))
    m = re.search(r"(\d+)\s*ft\s*[x×X]\s*(\d+)", p["Meta: spacing"])
    t = p["Meta: type"].strip()
    vals = {
        "SKU": p["SKU"], "Plant Name": p["Name"].strip(), "Old Category": oldcat, "New Category": newcat,
        "Sub-category": sub, "Show on site?": "Yes", "Best Seller?": "Yes" if p["Is featured?"] == "1" else "No",
        "Plant Type (Grafted / Seeded / Air layered...)": TYPEFIX.get(t.lower(), t),
        "Regular Price (Rs)": num(p["Regular price"]), "Sale Price (Rs)": num(p["Sale price"]),
        "Offline Price (Rs)": num(p["Meta: offlineprice"]), "Current Stock": num(p["Stock"]),
        "Availability": "Available" if p["In stock?"] == "1" else "Out of Stock",
        "Spacing L (ft)": int(m.group(1)) if m else None, "Spacing W (ft)": int(m.group(2)) if m else None,
        "Spacing (old text)": p["Meta: spacing"], "Exposure": p["Meta: exposure"], "Temperature": p["Meta: temperature"],
        "Growth Habit": p["Meta: growth-habit"], "Bearing Time": p["Meta: bearing-time"],
        "Yield per Tree": p["Meta: yield-per-tree"], "Harvest Season": p["Meta: harvesting-season"],
        "Fruit Size": p["Meta: fruit-size"], "Image URL (old site)": p["Images"].split(",")[0].strip(),
    }
    for k, v in vals.items():
        ws.cell(row=r, column=col[k], value=v if v != "" else None)
    L, W = get_column_letter(col["Spacing L (ft)"]), get_column_letter(col["Spacing W (ft)"])
    ws.cell(row=r, column=col["Plants per Acre"], value=f'=IF(AND(N({L}{r})>0,N({W}{r})>0),ROUND(43560/({L}{r}*{W}{r}),0),"")')
    r += 1
last = r - 1 + 30  # 30 blank rows for new plants
style_body(ws, 5, last, len(spec), fills)
for name in ("Regular Price (Rs)", "Sale Price (Rs)", "Offline Price (Rs)", "Bulk Price per Plant (Rs)"):
    for rr in range(5, last + 1):
        ws.cell(row=rr, column=col[name]).number_format = '"Rs "#,##0'
for rr in range(5, last + 1):
    ws.cell(row=rr, column=col["Plants per Acre"], value=ws.cell(row=rr, column=col["Plants per Acre"]).value or
            f'=IF(AND(N({L}{rr})>0,N({W}{rr})>0),ROUND(43560/({L}{rr}*{W}{rr}),0),"")')
CL = lambda n: get_column_letter(col[n])
dv_list(ws, ["Timber Trees", "Fruiting Trees", "Other Plants"], f"{CL('New Category')}5:{CL('New Category')}{last}")
dv_list(ws, ["Yes", "No"], f"{CL('Show on site?')}5:{CL('Best Seller?')}{last}")
dv_list(ws, ["Available", "Out of Stock", "Temporarily Unavailable", "Coming Soon", "Hidden"], f"{CL('Availability')}5:{CL('Availability')}{last}")
dv_list(ws, ["Main", "Host", "Biomass", "Main + Host", "Not used"], f"{CL('Role in Model (Main / Host / Biomass)')}5:{CL('Role in Model (Main / Host / Biomass)')}{last}")
dv_list(ws, ["Low", "Medium", "High"], f"{CL('Rain / Water Need')}5:{CL('Rain / Water Need')}{last}")
dv_list(ws, ["Yes", "No"], f"{CL('New Photo Needed?')}5:{CL('New Photo Needed?')}{last}")
dv_list(ws, ["0", "5", "12", "18"], f"{CL('GST %')}5:{CL('GST %')}{last}")
ws.auto_filter.ref = f"A4:{get_column_letter(len(spec))}{last}"
ws.freeze_panes = "C5"
n_plants = r - 5

# ---------- 4. Input Products ----------
ws = wb.create_sheet("Input Products")
title(ws, "Input Products", "Compost, vermicompost, biochar, seeds, liquid fertiliser, bio-pesticide. One row per pack size.")
cols = ["SKU", "Product Name", "Kannada Name", "Category", "Pack Size", "Unit (kg / L / g / packet)", "Pack Weight (kg)",
        "Regular Price (Rs)", "Offer Price (Rs)", "Current Stock", "Low Stock Alert at", "Availability", "HSN Code", "GST %",
        "Short Description", "How to Use / Dosage", "Recommended with which plants / models", "Best Seller?", "Image available?", "Notes"]
header(ws, 4, cols, [10, 24, 16, 16, 9, 10, 9, 10, 10, 9, 9, 14, 9, 7, 30, 30, 24, 9, 9, 22])
examples = ["Compost", "Vermicompost", "Biochar", "Seeds", "Liquid Fertiliser", "Bio-Pesticide"]
for i, cat in enumerate(examples):
    ws.cell(row=5 + i, column=4, value=cat)
style_body(ws, 5, 44, len(cols), [TOFILL] * len(cols))
dv_list(ws, examples + ["Other"], "D5:D44")
dv_list(ws, ["Available", "Out of Stock", "Temporarily Unavailable", "Coming Soon", "Hidden"], "L5:L44")
dv_list(ws, ["0", "5", "12", "18"], "N5:N44")
dv_list(ws, ["Yes", "No"], "R5:S44")

# ---------- 5. Recommendation Models ----------
ws = wb.create_sheet("Recommendation Models")
title(ws, "Recommendation Models (Help Me Choose)",
      "For each model, list the plants and their role. The site scales the 36 x 36 ft block to the customer's land size.")
r0 = 4
summary = ["Model", "Short description for customers", "Plant density (plants per 36 x 36 ft)", "Plants per acre",
           "Row pattern in 36 x 36 ft (e.g. M o X o X o M)", "Legend (what M / o / X mean)", "Min land size",
           "Planting video URL", "Notes"]
header(ws, r0, summary, [16, 34, 14, 12, 28, 26, 12, 26, 26])
for i, m in enumerate(["Commercial", "Food Forest", "Miyawaki"], 1):
    ws.cell(row=r0 + i, column=1, value=m)
style_body(ws, r0 + 1, r0 + 5, len(summary), [EXIST] + [TOFILL] * (len(summary) - 1))
r1 = r0 + 8
ws.cell(row=r1 - 1, column=1, value="Plants in each model").font = Font(name="Arial", bold=True, size=12, color=FOREST)
det = ["Model", "Plant Name (as in Plants tab)", "Category", "Role (Main / Host / Biomass)", "Layer",
       "Spacing L (ft)", "Spacing W (ft)", "Qty per 36 x 36 ft", "Qty per acre (auto)", "Swap options (alternative plants)",
       "Vermicompost per plant (kg)", "Biochar per plant (kg)", "Notes"]
for i, c in enumerate(det, 1):
    cell = ws.cell(row=r1, column=i, value=c)
    cell.fill, cell.font, cell.border = H_FILL, HFONT, BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[r1].height = 42
for w, i in zip([16, 26, 14, 12, 12, 9, 9, 10, 10, 26, 11, 10, 22], range(1, 14)):
    ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, w)
end = r1 + 60
for rr in range(r1 + 1, end + 1):
    # 36 x 36 ft = 1,296 sq ft; 43,560 / 1,296 = 33.61 blocks per acre
    ws.cell(row=rr, column=9, value=f'=IF(N(H{rr})>0,ROUND(H{rr}*43560/1296,0),"")')
style_body(ws, r1 + 1, end, len(det), [TOFILL] * 8 + [CALC] + [TOFILL] * 4)
for rr, m in zip(range(r1 + 1, r1 + 4), ["Commercial", "Food Forest", "Miyawaki"]):
    ws.cell(row=rr, column=1, value=m)
dv_list(ws, ["Commercial", "Food Forest", "Miyawaki"], f"A{r1+1}:A{end}")
dv_list(ws, ["Timber Trees", "Fruiting Trees", "Other Plants"], f"C{r1+1}:C{end}")
dv_list(ws, ["Main", "Host", "Biomass"], f"D{r1+1}:D{end}")
dv_list(ws, ["Canopy", "Sub-canopy", "Shrub", "Herb", "Ground cover", "Climber"], f"E{r1+1}:E{end}")
ws.freeze_panes = None

# ---------- 6. Calculators ----------
ws = wb.create_sheet("Calculators")
title(ws, "Calculator Formulas", "Our proposed formula is shown. Please confirm or replace it in the yellow column.")
cols = ["Calculator", "Where used", "Proposed formula (WODO)", "Inputs the customer enters", "Hasiru Agro formula / correction", "Example to test", "Approved?"]
header(ws, 4, cols, [24, 22, 46, 28, 46, 30, 10])
calcs = [
    ("Plants per acre", "Home, Shop (floating)", "Plants per acre = 43,560 / (spacing L ft x spacing W ft). Optional: + buffer % for replacements.", "Spacing L and W (ft), or plant picked (auto-fills its spacing), land size", "", "10 x 10 ft -> 436 plants/acre"),
    ("Buffer / mortality allowance", "Plant calculator", "Add X% extra plants for replacements (sketch shows '43,560 x 3 / 100'; is that a 3% allowance?).", "-", "", ""),
    ("Compost requirement", "Home compost calculator", "Compost kg = sum over plants (qty x compost kg per plant).", "Plants and quantities selected", "", "500 plants -> ? kg"),
    ("Vermicompost requirement", "Recommendation page, Cart", "Vermicompost kg = sum (qty x kg per plant, by plant size/type).", "Plants selected", "", "500 plants -> ? kg"),
    ("Biochar requirement", "Recommendation page, Cart", "Biochar kg = sum (qty x kg per plant).", "Plants selected", "", "500 plants -> ? kg"),
    ("Seed requirement", "Recommendation page", "Seeds (g or packets) per acre or per 36 x 36 block, e.g. for ground cover / green manure.", "Land size, model", "", ""),
    ("Model scaling", "Help Me Choose", "Plants for land = qty per 36 x 36 ft block x (land sq ft / 1,296).", "Land size (acre or sq ft)", "", "1 acre = 33.6 blocks"),
    ("Pack rounding", "Cart", "Input kg rounded up to whole packs (e.g. 112 kg -> 5 x 25 kg).", "-", "", ""),
]
for i, c in enumerate(calcs, 5):
    for j, v in enumerate(c + ("",), 1):
        ws.cell(row=i, column=j, value=v)
style_body(ws, 5, 4 + len(calcs), len(cols), [EXIST, EXIST, EXIST, EXIST, TOFILL, TOFILL, TOFILL])
dv_list(ws, ["Yes", "No", "Changed"], f"G5:G{4+len(calcs)}")

# ---------- 7. Logistics ----------
ws = wb.create_sheet("Logistics")
title(ws, "Delivery, Truck Capacity and Packing", "Needed for the truck-capacity meter, delivery charges and packing charges.")
ws.column_dimensions["A"].width = 34
for i, w in enumerate([16, 16, 16, 16, 16, 30], 2):
    ws.column_dimensions[get_column_letter(i)].width = w


def block(start, heading, cols, nrows, prefill=()):
    ws.cell(row=start, column=1, value=heading).font = Font(name="Arial", bold=True, size=12, color=FOREST)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=start + 1, column=i, value=c)
        cell.fill, cell.font, cell.border = H_FILL, HFONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[start + 1].height = 36
    for i, row in enumerate(prefill):
        for j, v in enumerate(row, 1):
            ws.cell(row=start + 2 + i, column=j, value=v)
    style_body(ws, start + 2, start + 1 + nrows, len(cols), [EXIST if prefill else TOFILL] + [TOFILL] * (len(cols) - 1))
    return start + nrows + 4


nxt = block(4, "A. Vehicles", ["Vehicle type", "Max weight (kg)", "Max plants (small saplings)", "Max plants (large / 6 ft+)", "Rate per km (Rs)", "Minimum charge (Rs)", "Notes"], 5,
            [("Tata Ace / small pickup",), ("Bolero pickup",), ("407 / mid truck",), ("Large truck",)])
nxt = block(nxt, "B. Bengaluru scheduled delivery", ["Setting", "Value", "", "", "", "", "Notes"], 6,
            [("Delivery day(s) of week",), ("Order cut-off (day + time)",), ("Delivery charge (flat / slab)",), ("Free delivery above (Rs)",), ("PIN codes / areas covered",), ("Max distance from nursery (km)",)])
nxt = block(nxt, "C. Outstation delivery", ["Rule", "Value", "", "", "", "", "Notes"], 7,
            [("Rate per km (Rs)",), ("Rate per kg (Rs)",), ("Minimum order for outstation (Rs / plants)",), ("States / districts served",), ("Delivery lead time (days)",), ("Shared truck / part-load allowed?",), ("Who pays tolls / loading?",)])
nxt = block(nxt, "D. Packing charges", ["Rule", "Value", "", "", "", "", "Notes"], 5,
            [("Packing type (per order / per plant / per kg / by size)",), ("Rate",), ("Free packing above (Rs)",), ("Special packing for large plants",), ("Courier option for single plants?",)])
block(nxt, "E. Nursery pickup", ["Setting", "Value", "", "", "", "", "Notes"], 4,
      [("Nursery address + Google Maps link",), ("Pickup days and timings",), ("Ready-for-pickup time after order",), ("Contact number for pickup",)])

# ---------- 8. Content & Policies ----------
ws = wb.create_sheet("Content & Policies")
title(ws, "Content, Media and Policies Checklist", "Share files in a shared Google Drive folder and mark the status here.")
cols = ["Item", "Used on", "Details needed", "Status", "Drive link / Notes"]
header(ws, 4, cols, [32, 22, 60, 14, 36])
items = [
    ("Home banner image + headline", "Home", "1-3 hero images (farm / nursery), headline idea."),
    ("Logos: UAS Dharwad, Startup India, MSME, others", "Home band, About", "High-res logos and approved wording ('Incubated by UAS Dharwad, recognised by Startup India')."),
    ("Proven results numbers", "Why Choose Us, Farm Dev", "Total acres developed, farms developed, clients, projects, plants planted (verified). Current site says 120+ farms, 850+ acres; please confirm."),
    ("Why Choose Us points", "Home, Why Choose Us", "Plant authenticity (how quality is ensured), domain expertise, delivery to farm."),
    ("About Us", "About", "Chemical-free approach, certified permaculture designer (certificate copy), vision, mission, team photos."),
    ("Testimonials", "Home", "Name, location, photo, review, rating (1-5), farm/project details. Minimum 6."),
    ("DIY / Learn-how videos", "Home, Recommendation", "YouTube links, title, category, which model it belongs to (Commercial / Food Forest / Miyawaki)."),
    ("Our Work (projects)", "Farm Development", "Per project: photos, before/after, location, acres, scope of work."),
    ("Our Clients", "Farm Development", "Client name (permission to show), project type, location, acreage, work done."),
    ("Damage policy", "Checkout, footer, policy page", "Reporting window (e.g. 48 hours), photo/video proof, what is covered."),
    ("Replacement policy", "Checkout, footer, policy page", "When a replacement is given, timeline, delivery of replacement."),
    ("Cancellation / return policy", "Checkout, footer, policy page", "Confirmed: no cancellation or return. Any exception (e.g. order not dispatched)?"),
    ("Shipping policy", "Policy page", "Dispatch time, delivery time, areas served."),
    ("Terms, Privacy policy", "Footer", "Company legal name, registered address, jurisdiction city."),
    ("GST / invoice details", "Invoices", "Company GSTIN, legal name, address, state code, invoice prefix."),
    ("Payment gateway account", "Checkout", "Razorpay / PhonePe / Cashfree, with KYC completed."),
    ("SMS / WhatsApp provider", "OTP, alerts", "Provider account, sender ID / DLT templates, WhatsApp Business number, admin alert numbers."),
    ("Kannada translations", "Whole site", "Who will translate/approve Kannada text and plant names?"),
    ("Current offers", "Cart", "Active and planned offers/coupons with dates and eligible products."),
    ("Contact details", "Footer, Contact", "Phone(s), email, WhatsApp, nursery address, map link, working hours."),
    ("Product photos", "Shop", "New photos for plants marked 'New Photo Needed?' = Yes, and all input products."),
]
for i, it in enumerate(items, 5):
    for j, v in enumerate(it + ("Pending",), 1):
        ws.cell(row=i, column=j, value=v)
style_body(ws, 5, 4 + len(items), len(cols), [EXIST, EXIST, EXIST, TOFILL, TOFILL])
dv_list(ws, ["Pending", "Shared", "Approved", "Not applicable"], f"D5:D{4+len(items)}")

# ---------- 9. Category Map ----------
ws = wb.create_sheet("Category Map")
title(ws, "Category Mapping (proposed)", "Old shop had 5 categories. The SRS asks for Timber / Fruiting / Other. Please approve or change.")
cols = ["Old category", "No. of plants", "Proposed new category", "Proposed filter / sub-category", "Approved? / Change to"]
header(ws, 4, cols, [20, 12, 22, 26, 30])
pc = get_column_letter(col["Old Category"])
for i, (old, (new, sub)) in enumerate(OLD2NEW.items(), 5):
    for j, v in enumerate((old, f"=COUNTIF(Plants!{pc}:{pc},A{i})", new, sub or "-", ""), 1):
        ws.cell(row=i, column=j, value=v)
style_body(ws, 5, 9, len(cols), [EXIST, CALC, EXIST, EXIST, TOFILL])
ws.cell(row=11, column=1, value="Input categories (new): Compost, Vermicompost, Biochar, Seeds, Liquid Fertiliser, Bio-Pesticide. Add more if needed.").font = BFONT

wb.save(OUT)
print("plants:", n_plants)
